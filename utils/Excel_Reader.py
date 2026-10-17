
from pathlib import Path
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = PROJECT_ROOT / "testData" / "InValidUsers.xlsx"


def get_excel_data(sheet_name):
    workbook = load_workbook(
        filename=EXCEL_PATH,
        data_only=True,
        read_only=True
    )

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)

        headers = next(rows, None)

        if headers is None:
            return []

        # Strip spaces from column names only.
        headers = [
            str(header).strip() if header is not None else ""
            for header in headers
        ]

        test_data = []

        for row in rows:
            if all(value is None for value in row):
                continue

            data = {
                header: "" if value is None else str(value)
                for header, value in zip(headers, row)
                if header
            }

            test_data.append(data)

        return test_data

    finally:
        workbook.close()
